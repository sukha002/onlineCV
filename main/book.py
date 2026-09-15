from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

# Base font
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(19)
styles["Heading 1"].font.bold = True
styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.size = Pt(14)
styles["Heading 2"].font.bold = True
# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.space_after = Pt(12)
r = p.add_run("Django Portfolio\nDeployment Handbook")
r.bold = True
r.font.size = Pt(30)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Local Django → GitHub → Azure App Service → Custom Domain")
r.italic = True
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\nA practical step-by-step record of the deployment process").font.size = Pt(11)

doc.add_page_break()

# Table of contents-like overview
doc.add_heading("Contents", level=1)
contents = [
    "1. Project Preparation",
    "2. GitHub Repository",
    "3. Resolving Git Synchronization",
    "4. Preparing Django for Azure",
    "5. Static Files and WhiteNoise",
    "6. Azure App Service",
    "7. Connecting GitHub to Azure",
    "8. Gunicorn Startup Configuration",
    "9. Testing and Troubleshooting",
    "10. Custom Domain and HTTPS",
    "11. Final Deployment Workflow",
    "12. Deployment Checklist",
]
for item in contents:
    doc.add_paragraph(item, style="List Bullet")

doc.add_page_break()

sections = [
("1. Project Preparation", [
"Build and test the Django portfolio locally.",
"Project structure used: onlineCV/ for the Django project and main/ for the application.",
"Templates are stored under main/templates/main/.",
"CSS and JavaScript are stored under main/static/main/css/ and main/static/main/js/.",
"Test the homepage at / and the About page at /about/ before deploying."
]),
("2. GitHub Repository", [
"GitHub repository: sukha002/onlineCV.",
"Use the main branch for deployment.",
"Keep the Django project, manage.py, requirements.txt, application code, templates, static files, and Azure workflow files synchronized with GitHub."
]),
("3. Resolving Git Synchronization", [
"If git push origin main is rejected with a non-fast-forward error, the remote branch contains commits that are not in the local branch.",
"To combine the histories safely, use: git pull origin main --no-rebase.",
"Resolve conflicts if Git reports them.",
"Commit the resolved changes and then use: git push origin main.",
"Do not use git push --force for this deployment workflow."
]),
("4. Preparing Django for Azure", [
"Install Gunicorn: pip install gunicorn.",
"Ensure gunicorn is included in requirements.txt.",
"The Azure startup module for this project is onlineCV.wsgi."
]),
("5. Static Files and WhiteNoise", [
"Azure initially loaded the Django HTML but the homepage appeared unstyled because /static/main/css/style.css returned 404.",
"Configure static files with STATIC_URL = '/static/' and STATIC_ROOT = BASE_DIR / 'staticfiles'.",
"Install WhiteNoise with: pip install whitenoise.",
"Add whitenoise to requirements.txt.",
"Add whitenoise.middleware.WhiteNoiseMiddleware immediately after django.middleware.security.SecurityMiddleware.",
"Run: python manage.py collectstatic --noinput.",
"The deployment successfully collected 129 static files.",
"Run: python manage.py check and confirm that Django reports no issues."
]),
("6. Azure App Service", [
"Use Azure App Service rather than Azure Static Web Apps for this Django server-side application.",
"Configure the Web App for Linux and Python.",
"Connect the App Service to the GitHub repository.",
"The Azure App Service default hostname used during testing was sukhwindersingh-gndadxcec2h0dbfr.canadacentral-01.azurewebsites.net."
]),
("7. Connecting GitHub to Azure", [
"Open App Service → Deployment Center.",
"Select GitHub as the deployment source.",
"Select repository sukha002/onlineCV and branch main.",
"Use GitHub Actions for deployment.",
"Check Deployment Center logs and confirm the deployment status is Success."
]),
("8. Gunicorn Startup Configuration", [
"Azure initially displayed the default Python developer page instead of the Django portfolio.",
"Configure App Service → Configuration → General settings → Startup Command.",
"Use exactly: gunicorn --bind=0.0.0.0 --timeout 600 onlineCV.wsgi.",
"Save/apply the configuration and allow Azure to restart the application.",
"After restart, test the Azure default URL and /about/."
]),
("9. Testing and Troubleshooting", [
"Test the homepage and About page separately.",
"If HTML loads but styling is missing, test /static/main/css/style.css directly.",
"If the CSS URL returns 404, inspect STATIC_URL, STATIC_ROOT, collectstatic, and WhiteNoise configuration.",
"If Azure reports DisallowedHost, add the Azure hostname to ALLOWED_HOSTS.",
"Keep generated __pycache__ files out of Git commits.",
"Use Azure App Service → Monitoring → Log stream when the application returns a server error."
]),
("10. Custom Domain and HTTPS", [
"Only move the custom domain after the Azure default hostname works correctly.",
"Add sukhwindertoor.com and www.sukhwindertoor.com under Azure App Service → Custom domains.",
"Azure will provide the DNS verification records required for the domain.",
"Update the GoDaddy DNS records using the exact values supplied by Azure.",
"Do not remove or replace DNS records blindly; verify each record before changing it.",
"After domain verification, configure HTTPS/TLS and enable HTTPS Only."
]),
("11. Final Deployment Workflow", [
"Make changes locally.",
"Test Django locally.",
"Run python manage.py check.",
"Run collectstatic when static assets change.",
"Commit the changes.",
"Push to GitHub main.",
"GitHub Actions deploys the new version to Azure App Service.",
"Verify the production website and important routes after deployment."
]),
("12. Deployment Checklist", [
"☐ Django homepage works locally.",
"☐ /about/ works locally.",
"☐ CSS and JavaScript work locally.",
"☐ requirements.txt contains Django, Gunicorn, and WhiteNoise.",
"☐ STATIC_URL is defined once.",
"☐ STATIC_ROOT is configured.",
"☐ WhiteNoise middleware is configured.",
"☐ collectstatic completes successfully.",
"☐ python manage.py check reports no issues.",
"☐ GitHub main branch is synchronized.",
"☐ GitHub Actions deployment succeeds.",
"☐ Gunicorn startup command is configured.",
"☐ Azure homepage works.",
"☐ Azure /about/ works.",
"☐ Azure static files work.",
"☐ Custom domain is verified.",
"☐ HTTPS works.",
"☐ Final production website is tested."
]),
]

for title, bullets in sections:
    doc.add_heading(title, level=1)
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    doc.add_paragraph()

# Quick reference
doc.add_page_break()
doc.add_heading("Quick Command Reference", level=1)
commands = [
"python manage.py runserver",
"python manage.py check",
"python manage.py collectstatic --noinput",
"pip install gunicorn",
"pip install whitenoise",
"pip freeze > requirements.txt",
"git status",
"git add .",
"git commit -m \"Update portfolio\"",
"git pull origin main --no-rebase",
"git push origin main",
]
for c in commands:
    p = doc.add_paragraph()
    r = p.add_run(c)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

doc.add_heading("Final Architecture", level=1)
for line in [
    "Your Computer",
    "↓",
    "Django Project",
    "↓",
    "GitHub — sukha002/onlineCV",
    "↓",
    "GitHub Actions",
    "↓",
    "Azure App Service",
    "↓",
    "Gunicorn",
    "↓",
    "Django",
    "↓",
    "sukhwindertoor.com",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)

from pathlib import Path

path = Path(__file__).resolve().parent / "Django_Portfolio_Azure_Deployment_Handbook.docx"

doc.save(path)

print(f"File created successfully:")
print(path)
